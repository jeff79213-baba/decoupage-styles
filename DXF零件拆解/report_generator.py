import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, size=11, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
ISSUE_FILL = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
ISSUE_FONT = Font(color='9C0006')
BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

def generate_excel(analysis_result, output_path):
    wb = openpyxl.Workbook()
    ws_parts = wb.active
    ws_parts.title = '零件材料清單'
    headers = ['件號', '圖層', '形狀', '長度(mm)', '寬度(mm)', '厚度(mm)', '標註數', '備註']
    for col, h in enumerate(headers, 1):
        cell = ws_parts.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    for row_idx, part in enumerate(analysis_result.get('parts', []), 2):
        values = [
            part.get('part_id', ''),
            part.get('layer', ''),
            part.get('shape', ''),
            part.get('length'),
            part.get('width'),
            part.get('thickness'),
            part.get('dimension_count', 0),
            part.get('note', ''),
        ]
        for col, val in enumerate(values, 1):
            cell = ws_parts.cell(row=row_idx, column=col, value=val)
            cell.border = BORDER
            cell.alignment = CENTER if col <= 7 else LEFT
    for col in range(1, len(headers) + 1):
        ws_parts.column_dimensions[get_column_letter(col)].width = 14
    ws_issues = wb.create_sheet('問題清單')
    issue_headers = ['零件編號', '區域', '問題描述', '可能原因']
    for col, h in enumerate(issue_headers, 1):
        cell = ws_issues.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        cell.alignment = CENTER
        cell.border = BORDER
    for row_idx, issue in enumerate(analysis_result.get('issues', []), 2):
        values = [
            issue.get('part_id', ''),
            issue.get('zone', ''),
            issue.get('description', ''),
            issue.get('reason', ''),
        ]
        for col, val in enumerate(values, 1):
            cell = ws_issues.cell(row=row_idx, column=col, value=val)
            cell.fill = ISSUE_FILL
            cell.font = ISSUE_FONT
            cell.border = BORDER
            cell.alignment = LEFT
    for col in range(1, len(issue_headers) + 1):
        ws_issues.column_dimensions[get_column_letter(col)].width = 30
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else '.', exist_ok=True)
    wb.save(output_path)
    return output_path
